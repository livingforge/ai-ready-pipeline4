"""Pure planning for structural Excel edits; all coordinates are original-sheet coordinates.

Original cells follow row changes; inserted cells refer to an insertion ID + offset.
Explicit formulas and shape geometry describe the final workbook. No LLM inference here.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from .contracts import DocumentError, digest, under

MAX_ROW = 1048576


def column_index(column: str) -> int:
    if not re.fullmatch(r"[A-Z]+", column):
        raise DocumentError(f"invalid column: {column}")
    value = 0
    for char in column:
        value = value * 26 + ord(char) - ord("A") + 1
    if value > 16384:
        raise DocumentError(f"column exceeds Excel limit: {column}")
    return value


def coordinate(address: str) -> tuple[str, int]:
    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", address)
    if not match or int(match[2]) > MAX_ROW:
        raise DocumentError(f"invalid cell address: {address}")
    column_index(match[1])
    return match[1], int(match[2])


def validate_rows(operations: list[dict], sheets: set[str]) -> list[dict]:
    ids = [o["id"] for o in operations]
    if len(set(ids)) != len(ids):
        raise DocumentError("duplicate operation ID")
    rows = [o for o in operations if o["kind"] in ("insert_rows", "delete_rows")]
    for op in operations:
        if op["sheet"] not in sheets or not op["reason"].strip():
            raise DocumentError(f"operation needs existing sheet and reason: {op['id']}")
    for i, op in enumerate(rows):
        if op["at"] + op["count"] - 1 > MAX_ROW:
            raise DocumentError(f"row operation exceeds Excel limits: {op['id']}")
        if op["kind"] == "delete_rows" and "style_from" in op:
            raise DocumentError("style_from applies only to inserted rows")
        for other in rows[i + 1:]:
            if op["sheet"] != other["sheet"]:
                continue
            end = op["at"] + op["count"] - 1 if op["kind"] == "delete_rows" else op["at"]
            other_end = other["at"] + other["count"] - 1 if other["kind"] == "delete_rows" else other["at"]
            if max(op["at"], other["at"]) <= min(end, other_end):
                raise DocumentError(f"overlapping row operations: {op['id']}, {other['id']}")
        if "style_from" in op and transform_row(op["sheet"], op["style_from"], rows) is None:
            raise DocumentError("style template row is deleted")
    return sorted(rows, key=lambda o: (o["sheet"], -o["at"]))


def transform_row(sheet: str, row: int, rows: list[dict]) -> int | None:
    delta = 0
    for op in rows:
        if op["sheet"] != sheet:
            continue
        if op["kind"] == "delete_rows":
            if op["at"] <= row < op["at"] + op["count"]:
                return None
            if row >= op["at"] + op["count"]:
                delta -= op["count"]
        elif row >= op["at"]:
            delta += op["count"]
    result = row + delta
    if not 1 <= result <= MAX_ROW:
        raise DocumentError("row shift would exceed Excel limits")
    return result


def resolve_target(target: dict, rows: list[dict]) -> dict:
    if "cell" in target:
        column, row = coordinate(target["cell"])
        final = transform_row(target["sheet"], row, rows)
        if final is None:
            raise DocumentError(f"writeback targets a deleted row: {target}")
        return {"sheet": target["sheet"], "cell": f"{column}{final}"}
    matches = [o for o in rows if o["id"] == target["insertion"] and o["kind"] == "insert_rows" and o["sheet"] == target["sheet"]]
    if not matches or not 0 <= target["offset"] < matches[0]["count"]:
        raise DocumentError(f"invalid insertion target: {target}")
    op = matches[0]
    column_index(target["column"])
    prior = [other for other in rows if other["sheet"] == op["sheet"] and other["at"] < op["at"]]
    delta = sum(other["count"] if other["kind"] == "insert_rows" else -other["count"] for other in prior)
    final = op["at"] + delta + target["offset"]
    if not 1 <= final <= MAX_ROW:
        raise DocumentError("inserted cell exceeds Excel limits")
    return {"sheet": target["sheet"], "cell": f"{target['column']}{final}"}


def formula(value) -> None:
    if value is None:
        return  # Explicitly clear a formula.
    if not isinstance(value, str) or not value.startswith("=") or not value[1:].strip():
        raise DocumentError("formula writeback requires an =formula string or null")
    if re.search(r"[\x00-\x1f]", value):
        raise DocumentError("control characters in formula")
    from openpyxl.formula import Tokenizer
    try:
        Tokenizer(value)
    except Exception as exc:
        raise DocumentError(f"invalid formula: {exc}") from exc


def resolve_operations(operations: list[dict], values: dict, mappings: list[dict],
                       directory: Path, drawings: list[dict]) -> list[dict]:
    """Resolve shape properties from the canonical Markdown, never duplicate their values."""
    active = {(m["page"], m["block"], m["field"]): m for m in mappings if m["writeback"] == "operation"}
    used = set()
    def field(ref):
        key = (ref["page"], ref["block"], ref["field"])
        if key not in values or key not in active:
            raise DocumentError(f"operation field requires an operation mapping: {key}")
        used.add(key)
        return values[key]
    names = {(s["sheet"], s["name"]): s for s in drawings}
    if len(names) != len(drawings):
        raise DocumentError("ambiguous drawing names in original workbook")
    modified, added = set(), set()
    resolved = []
    for op in operations:
        if op["kind"] in ("insert_rows", "delete_rows"):
            resolved.append(dict(op))
            continue
        key = (op["sheet"], op["name"])
        if key in modified:
            raise DocumentError(f"multiple operations on the same shape: {key}")
        modified.add(key)
        is_new = op["kind"] in ("add_shape", "add_picture", "add_connector")
        if is_new:
            if key in names or key in added:
                raise DocumentError(f"shape already exists: {key}")
            added.add(key)
        elif key not in names:
            raise DocumentError(f"shape not found: {key}; use documents drawings")
        if op["kind"] == "replace_picture" and names[key]["kind"] != "pic":
            raise DocumentError("replace_picture target is not a picture")
        result = {**op, "properties": {name: field(ref) for name, ref in op.get("properties", {}).items()}}
        properties = result["properties"]
        for name, value in properties.items():
            if name in ("text", "fill", "line"):
                if not isinstance(value, str):
                    raise DocumentError(f"shape {name} must be a string")
                if name in ("fill", "line") and not re.fullmatch(r"#[0-9A-Fa-f]{6}|none", value):
                    raise DocumentError("shape colors must be #RRGGBB or none")
            else:
                if type(value) not in (int, float) or not math.isfinite(value):
                    raise DocumentError(f"shape {name} must be a finite number")
                line_shape = op.get("shape_type") == "line" or op["kind"] == "add_connector" or names.get(key, {}).get("kind") == "cxnSp"
                positive = name in ("font_size", "line_weight") or name in ("width", "height") and not line_shape
                if value < 0 or positive and value == 0:
                    raise DocumentError(f"invalid shape {name}: {value}")
                if name == "rotation" and value >= 360:
                    raise DocumentError("rotation must be in [0, 360)")
        if op["kind"] in ("add_shape", "add_picture") and not {"left", "top", "width", "height"} <= properties.keys():
            raise DocumentError("new shapes require left/top/width/height in points")
        if op.get("shape_type") == "line" and properties.get("width") == properties.get("height") == 0:
            raise DocumentError("a line needs nonzero length")
        if op["kind"] in ("add_picture", "replace_picture"):
            asset = field(op["asset"])
            if not isinstance(asset, str):
                raise DocumentError("picture asset must be a path string")
            path = under(directory, asset)
            if not path.is_relative_to(directory.resolve() / "assets") or not path.is_file():
                raise DocumentError("picture must exist inside the document assets directory")
            from PIL import Image
            try:
                with Image.open(path) as image:
                    image.verify()
            except Exception as exc:
                raise DocumentError(f"invalid image: {asset}") from exc
            result["asset"] = asset
            result["asset_sha256"] = digest(path.read_bytes())
        if op["kind"] in ("add_picture", "replace_picture", "add_connector") and any(k in properties for k in ("text", "font_size")):
            raise DocumentError("pictures/connectors cannot carry text properties")
        resolved.append(result)
    deleted = {(o["sheet"], o["name"]) for o in operations if o["kind"] == "delete_shape"}
    for op in operations:
        if op["kind"] == "add_connector":
            for endpoint in (op["begin"], op["end"]):
                key = (op["sheet"], endpoint["name"])
                if key in deleted or key not in names and key not in added or key == (op["sheet"], op["name"]):
                    raise DocumentError(f"connector endpoint not found: {key}")
    if used != set(active):
        raise DocumentError(f"operation mappings not consumed by an operation: {set(active) - used}")
    return resolved
