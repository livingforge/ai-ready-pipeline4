"""Structural writeback contracts and real Excel integration (ARP_TEST_EXCEL=1)."""
import os
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from arp4.documents import excel, markdown, native_excel, operations
from arp4.documents.contracts import DocumentError, read, validate, write
from arp4.documents.store import Store


def row(kind, at, count=1, name="rows"):
    return {"id": name, "kind": kind, "sheet": "Data", "at": at, "count": count, "reason": "承認済みの項目変更"}


def test_resize_preserves_aspect_ratio_setting_even_on_failure():
    class Shape:
        LockAspectRatio = -1
        Width = 10

        @property
        def Height(self):
            return 10

        @Height.setter
        def Height(self, value):
            raise RuntimeError("resize failed")

    shape = Shape()
    native_excel._properties(shape, {"width": 20})
    assert shape.Width == 20
    assert shape.LockAspectRatio == -1
    with pytest.raises(RuntimeError, match="resize failed"):
        native_excel._properties(shape, {"height": 20})
    assert shape.LockAspectRatio == -1


def test_original_coordinates_and_multiple_insertions():
    rows = operations.validate_rows([row("insert_rows", 2, 2, "new"), row("delete_rows", 6, 2, "remove")], {"Data"})
    assert operations.resolve_target({"sheet": "Data", "cell": "B3"}, rows)["cell"] == "B5"
    assert operations.resolve_target({"sheet": "Data", "cell": "B8"}, rows)["cell"] == "B8"
    assert operations.resolve_target({"sheet": "Data", "insertion": "new", "offset": 1, "column": "B"}, rows)["cell"] == "B3"
    with pytest.raises(DocumentError, match="deleted"):
        operations.resolve_target({"sheet": "Data", "cell": "B6"}, rows)
    with pytest.raises(DocumentError, match="overlapping"):
        operations.validate_rows([row("insert_rows", 4, name="a"), row("delete_rows", 3, 3, "b")], {"Data"})
    last = operations.validate_rows([row("insert_rows", 1048576)], {"Data"})
    assert operations.resolve_target({"sheet": "Data", "insertion": "rows", "offset": 0, "column": "XFD"}, last)["cell"] == "XFD1048576"


def test_missing_insertion_or_invalid_formula_is_not_silently_accepted():
    with pytest.raises(DocumentError, match="invalid insertion"):
        operations.resolve_target({"sheet": "Data", "insertion": "missing", "offset": 0, "column": "A"}, [])
    for value in ("SUM(A1:A2)", "=", 12, "=A1\x00"):
        with pytest.raises(DocumentError):
            operations.formula(value)
    operations.formula("=SUM(A1:A2)")


def add_value(directory, mappings, key, value, *, mode="operation", target=None):
    path = directory / "content/changes.md"
    if not path.exists():
        meta = read(directory / "document.yml")
        path.write_text('---\nschema_version: "1"\ndocument_id: ' + meta["document_id"]
            + '\npage_id: changes\n---\n\n# 変更\n\n<!-- arp:block id="changes" -->\n\n'
            + '| id | type | value |\n| --- | --- | --- |\n', encoding="utf-8")
        mappings["entries"].append({"page": "changes", "block": "changes", "field": None,
            "origins": [], "reason": "承認済み変更の一覧", "writeback": "excluded", "target": None})
    with path.open("a", encoding="utf-8") as stream:
        stream.write(markdown.value_row(key, value) + "\n")
    mappings["entries"].append({"page": "changes", "block": "changes", "field": key,
        "origins": [], "reason": "承認済み変更", "writeback": mode, "target": target})
    return {"page": "changes", "block": "changes", "field": key}


def adopt(store, proposal):
    mappings = read(proposal / "mappings.yml")
    for mapping in mappings["entries"]:
        if mapping["writeback"] == "pending":
            mapping.update(writeback="excluded", reason="説明文。値は型付き項目から反映")
    write(proposal / "mappings.yml", mappings)
    prompt = store.root / "prompt.md"
    prompt.write_text("承認済み変更の成形", encoding="utf-8")
    store.record(proposal.name, "human", "editor", prompt)
    return store.adopt(proposal.name, "reviewer")


@pytest.fixture
def source(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    for line in (("Name", "Amount"), ("First", 10), ("Second", 20), ("Third", 30)):
        sheet.append(line)
    sheet["C2"] = "=B2*2"
    sheet["B2"].number_format = "0.00"
    sheet.add_table(Table(displayName="Orders", ref="A1:B4"))
    book.defined_names.add(DefinedName("Measure", attr_text="'Data'!$B$2"))
    book.create_sheet("Other")["A1"] = "=Data!B2"
    path = tmp_path / "source.xlsx"
    book.save(path)
    return path


def test_v1_read_compatibility_and_v2_shape_fields(source):
    store = Store.init(source.parent)
    proposal = store.import_source(source, "doc")
    mappings = read(proposal / "mappings.yml")
    assert mappings["schema_version"] == "2"
    old = {**mappings, "schema_version": "1"}
    del old["operations"]
    validate("mappings", old)
    text = add_value(proposal, mappings, "text", "新規処理")
    geom = {key: add_value(proposal, mappings, key, value) for key, value in {"left": 20, "top": 30, "width": 120, "height": 60}.items()}
    mappings["operations"] = [{"id": "new-shape", "kind": "add_shape", "sheet": "Data", "name": "New",
        "shape_type": "rectangle", "properties": {**geom, "text": text}, "reason": "新規業務"}]
    write(proposal / "mappings.yml", mappings)
    assert store.inspect(proposal)["operations"][0]["properties"]["text"] == "新規処理"
    adopt(store, proposal)
    plan = store.export("doc")
    assert plan["engine"] == "excel"
    with pytest.raises(DocumentError, match="require the Excel"):
        store.export("doc", engine="xml")


def test_rows_cover_deleted_source_cells_and_reject_surviving_targets(source):
    store = Store.init(source.parent)
    proposal = store.import_source(source, "doc")
    mappings = read(proposal / "mappings.yml")
    mappings["operations"] = [row("delete_rows", 4)]
    write(proposal / "mappings.yml", mappings)
    with pytest.raises(DocumentError, match="deleted row"):
        store.inspect(proposal)
    for mapping in mappings["entries"]:
        if mapping["target"] and mapping["target"]["sheet"] == "Data" and mapping["target"]["cell"].endswith("4"):
            mapping.update(writeback="excluded", reason="行削除で取り除く値")
    write(proposal / "mappings.yml", mappings)
    assert ("Data", "B4") in store.inspect(proposal)["deleted_cells"]


@pytest.fixture
def real_excel():
    if os.environ.get("ARP_TEST_EXCEL") != "1":
        pytest.skip("set ARP_TEST_EXCEL=1 on a machine with Microsoft Excel")
    return native_excel.application


def test_real_excel_rows_formulas_names_tables_and_literal_values(source, real_excel):
    store = Store.init(source.parent)
    proposal = store.import_source(source, "doc")
    mappings = read(proposal / "mappings.yml")
    mappings["operations"] = [row("insert_rows", 2, name="new"), row("delete_rows", 4, name="remove")]
    mappings["operations"][0]["style_from"] = 2
    for mapping in mappings["entries"]:
        if mapping["target"] and mapping["target"]["sheet"] == "Data" and mapping["target"]["cell"].endswith("4"):
            mapping.update(writeback="excluded", reason="行削除で取り除く値")
        if mapping["target"] == {"sheet": "Data", "cell": "C2"}:
            mapping.update(writeback="excluded", reason="変更一覧で数式を明示する")
    add_value(proposal, mappings, "label", "=literal text", mode="cell",
        target={"sheet": "Data", "insertion": "new", "offset": 0, "column": "A"})
    add_value(proposal, mappings, "amount", 40, mode="cell",
        target={"sheet": "Data", "insertion": "new", "offset": 0, "column": "B"})
    add_value(proposal, mappings, "formula", "=SUM(B2:B4)", mode="formula", target={"sheet": "Data", "cell": "C2"})
    write(proposal / "mappings.yml", mappings)
    adopt(store, proposal)
    output = source.parent / ".arp/out/rows.xlsx"
    report = store.export("doc", output)
    assert report["engine"] == "excel" and report["verified_after_reopen"]
    result = load_workbook(output)
    sheet = result["Data"]
    assert sheet["A2"].value == "=literal text" and sheet["A2"].data_type == "s"
    assert [sheet[f"B{i}"].value for i in range(2, 5)] == [40, 10, 20]
    assert sheet["B2"].number_format == "0.00"
    assert sheet["C3"].value == "=SUM(B2:B4)"
    assert result["Other"]["A1"].value == "=Data!B3"
    assert "$B$3" in result.defined_names["Measure"].attr_text
    assert sheet.tables["Orders"].ref == "A1:B4"
    assert load_workbook(source)["Data"]["B2"].value == 10


def test_real_excel_shape_edit_add_delete_connect_and_picture(source, real_excel):
    from PIL import Image
    shaped = source.with_name("with-shapes.xlsx")
    native_excel.patch(source, shaped, [], [
        {"id": "process", "kind": "add_shape", "sheet": "Data", "name": "Process", "shape_type": "rectangle",
         "properties": {"left": 20, "top": 100, "width": 100, "height": 40, "text": "Before"}, "reason": "fixture"},
        {"id": "obsolete", "kind": "add_shape", "sheet": "Data", "name": "Obsolete", "shape_type": "rectangle",
         "properties": {"left": 250, "top": 100, "width": 100, "height": 40}, "reason": "fixture"},
    ], source.parent)
    shaped.replace(source)
    assert {d["name"] for d in excel.drawings(source)} >= {"Process", "Obsolete"}
    store = Store.init(source.parent)
    proposal = store.import_source(source, "doc")
    mappings = read(proposal / "mappings.yml")
    props = {key: add_value(proposal, mappings, key, value) for key, value in {
        "text": "受注確認", "left": 30, "top": 120, "width": 140, "height": 50,
        "fill": "#00AA66", "line": "#112233", "line_weight": 2, "font_size": 14}.items()}
    new_props = {key: add_value(proposal, mappings, "new-" + key, value) for key, value in {
        "text": "完了", "left": 240, "top": 120, "width": 120, "height": 50}.items()}
    (proposal / "assets").mkdir(exist_ok=True)
    Image.new("RGB", (30, 20), "red").save(proposal / "assets/logo.png")
    asset = add_value(proposal, mappings, "logo", "assets/logo.png")
    picture_props = {key: add_value(proposal, mappings, "image-" + key, value) for key, value in {
        "left": 20, "top": 220, "width": 90, "height": 60}.items()}
    mappings["operations"] = [
        {"id": "update", "kind": "update_shape", "sheet": "Data", "name": "Process", "properties": props, "reason": "処理名変更"},
        {"id": "delete", "kind": "delete_shape", "sheet": "Data", "name": "Obsolete", "reason": "旧処理廃止"},
        {"id": "new", "kind": "add_shape", "sheet": "Data", "name": "Done", "shape_type": "rounded_rectangle", "properties": new_props, "reason": "完了処理"},
        {"id": "link", "kind": "add_connector", "sheet": "Data", "name": "Flow", "begin": {"name": "Process", "site": 4},
         "end": {"name": "Done", "site": 2}, "connector_type": "elbow", "reason": "処理順"},
        {"id": "logo", "kind": "add_picture", "sheet": "Data", "name": "Logo", "asset": asset, "properties": picture_props, "reason": "ロゴ追加"}]
    write(proposal / "mappings.yml", mappings)
    adopt(store, proposal)
    output = source.parent / ".arp/out/shapes.xlsx"
    report = store.export("doc", output)
    assert report["verified_after_reopen"]
    drawings = {d["name"]: d for d in excel.drawings(output)}
    assert drawings["Process"]["text"] == "受注確認"
    assert drawings["Done"]["text"] == "完了"
    assert drawings["Logo"]["kind"] == "pic"
    assert drawings["Flow"]["kind"] == "cxnSp"
    assert "Obsolete" not in drawings
    assert "Obsolete" in {d["name"] for d in excel.drawings(source)}
    # Re-import the generated workbook and replace a picture without changing its geometry.
    incoming = source.parent / "returned.xlsx"
    incoming.write_bytes(output.read_bytes())
    next_proposal = store.import_source(incoming, "returned")
    replacement = read(next_proposal / "mappings.yml")
    (next_proposal / "assets").mkdir(exist_ok=True)
    Image.new("RGB", (30, 20), "blue").save(next_proposal / "assets/new-logo.png")
    asset = add_value(next_proposal, replacement, "new-logo", "assets/new-logo.png")
    replacement["operations"] = [{"id": "replace", "kind": "replace_picture", "sheet": "Data", "name": "Logo",
                                  "asset": asset, "properties": {}, "reason": "ロゴ更新"}]
    write(next_proposal / "mappings.yml", replacement)
    adopt(store, next_proposal)
    updated = source.parent / ".arp/out/replaced.xlsx"
    assert store.export("returned", updated)["verified_after_reopen"]
    from io import BytesIO
    with ZipFile(updated) as archive:
        images = [Image.open(BytesIO(archive.read(name))).convert("RGB") for name in archive.namelist()
                  if name.startswith("xl/media/") and name.endswith(".png")]
        assert any(image.getpixel((0, 0)) == (0, 0, 255) for image in images)


def test_real_excel_invalid_formula_leaves_no_output(source, real_excel):
    store = Store.init(source.parent)
    proposal = store.import_source(source, "doc")
    mappings = read(proposal / "mappings.yml")
    add_value(proposal, mappings, "broken-formula", "=1+", mode="formula", target={"sheet": "Data", "cell": "C2"})
    write(proposal / "mappings.yml", mappings)
    adopt(store, proposal)
    before = source.read_bytes()
    output = source.parent / ".arp/out/broken.xlsx"
    with pytest.raises(DocumentError):
        store.export("doc", output)
    assert not output.exists() and not output.with_suffix(".xlsx.report.json").exists()
    assert source.read_bytes() == before
