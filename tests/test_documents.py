"""End-to-end authority editing, provenance, conflicts and byte-preserving Excel updates."""
from pathlib import Path
import shutil
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from arp4 import cli
from arp4.documents import excel, markdown
from arp4.documents.contracts import DocumentError, read, write
from arp4.documents.store import Store


@pytest.fixture
def project(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "項目定義"
    sheet.append(["受注番号", 10, True])
    sheet.append(["備考", "初期値", "=B1*2"])
    sheet["B1"].number_format = "0000"
    source = tmp_path / "基本設計.xlsx"
    book.save(source)
    store = Store.init(tmp_path)
    proposal = store.import_source(source, "order-design")
    return store, source, proposal


def ready(store, proposal):
    mappings = read(proposal / "mappings.yml")
    for entry in mappings["entries"]:
        if entry["writeback"] == "pending":
            entry.update(writeback="excluded", reason="説明部分。変更は型付きセル表から反映")
    write(proposal / "mappings.yml", mappings)
    prompt = store.root / "prompt.txt"
    prompt.write_text("原文の意味を変えずに成形し、出典とセル対応を維持する", encoding="utf-8")
    store.record(proposal.name, "test-model", "test-agent", prompt)
    return store.adopt(proposal.name, "reviewer")


def test_excel_roundtrip_keeps_unrelated_parts_and_refuses_unreviewed(project):
    store, source, proposal = project
    directory = ready(store, proposal)
    page = directory / "content/sheet-1.md"
    text = page.read_text(encoding="utf-8").replace("| c-1-B1 | number | 10 |", "| c-1-B1 | number | 12 |")
    page.write_text(text, encoding="utf-8")
    assert not store.inspect(directory)["reviewed"]
    with pytest.raises(DocumentError, match="not reviewed"):
        store.export("order-design")
    store.review("order-design", "reviewer")
    plan = store.export("order-design")
    assert plan["changes"] == [{"sheet": "項目定義", "cell": "B1", "before": 10, "after": 12, "field": "c-1-B1"}]
    output = store.root / ".arp/out/revised.xlsx"
    report = store.export("order-design", output)
    assert report["written"] and report["requires_excel_recalculation"]
    actual = load_workbook(output)
    assert actual["項目定義"]["B1"].value == 12
    assert actual["項目定義"]["B1"].number_format == "0000"
    assert actual["項目定義"]["C2"].value == "=B1*2"
    with ZipFile(source) as old, ZipFile(output) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name not in report["changed_parts"]:
                assert old.read(name) == new.read(name)
    assert load_workbook(source)["項目定義"]["B1"].value == 10
    with pytest.raises(DocumentError, match="already exists"):
        store.export("order-design", output)


def test_original_drift_blocks_export_and_reimport_protects_edits(project):
    store, source, proposal = project
    directory = ready(store, proposal)
    source = store.root / read(directory / "document.yml")["source"]["path"]
    original = (directory / "content/sheet-1.md").read_bytes()
    book = load_workbook(source)
    book["項目定義"]["B1"] = 20
    book.save(source)
    assert not store.inspect(directory)["source_current"]
    with pytest.raises(DocumentError, match="source changed"):
        store.export("order-design")
    incoming = store.import_source(source, "order-design")
    assert (directory / "content/sheet-1.md").read_bytes() == original
    page = directory / "content/sheet-1.md"
    page.write_text(page.read_text(encoding="utf-8").replace("10 |", "11 |"), encoding="utf-8")
    with pytest.raises(DocumentError, match="authority changed"):
        ready(store, incoming)
    assert store.diff(incoming.name)["authority_changed"]


def test_deleted_marker_and_duplicate_or_unmapped_fields_are_errors(project):
    store, _, proposal = project
    page = proposal / "content/sheet-1.md"
    text = page.read_text(encoding="utf-8")
    page.write_text(text.replace('<!-- arp:block id="cells" -->', ""), encoding="utf-8")
    with pytest.raises(DocumentError, match="must belong"):
        store.inspect(proposal)
    page.write_text(text + "| c-1-B1 | number | 17 |\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="duplicate field"):
        store.inspect(proposal)
    page.write_text(text + "| new-id | number | 17 |\n", encoding="utf-8")
    with pytest.raises(DocumentError, match="coverage"):
        store.inspect(proposal)


def test_fenced_marker_is_literal_and_malformed_marker_rejected(tmp_path):
    page = tmp_path / "page.md"
    text = '---\nschema_version: "1"\ndocument_id: doc\npage_id: p\n---\n\n# Title\n\n<!-- arp:block id="real" -->\n\n```html\n<!-- arp:block id="fake" -->\n```\n'
    page.write_text(text, encoding="utf-8")
    assert set(markdown.parse(page, "doc").blocks) == {"real"}
    page.write_text(text + '\n<!-- arp:block id=broken -->\n', encoding="utf-8")
    with pytest.raises(DocumentError, match="standalone"):
        markdown.parse(page, "doc")


def test_bad_table_width_and_type(project):
    store, _, proposal = project
    page = proposal / "content/sheet-1.md"
    text = page.read_text(encoding="utf-8")
    page.write_text(text.replace("| c-1-B1 | number | 10 |", "| c-1-B1 | number | 10 | extra |"), encoding="utf-8")
    with pytest.raises(DocumentError, match="column count"):
        store.inspect(proposal)
    page.write_text(text.replace("| c-1-B1 | number | 10 |", '| c-1-B1 | number | "ten" |'), encoding="utf-8")
    with pytest.raises(DocumentError, match="type/value"):
        store.inspect(proposal)


def test_formula_target_and_duplicate_yaml_rejected(project):
    store, _, proposal = project
    mappings = read(proposal / "mappings.yml")
    for entry in mappings["entries"]:
        if entry["field"] == "c-1-C2":
            entry["writeback"] = "cell"
    write(proposal / "mappings.yml", mappings)
    with pytest.raises(DocumentError, match="formula/error"):
        store.inspect(proposal)
    config = store.arp / "config.yml"
    config.write_text('schema_version: "1"\ndirectory: knowledge\ndirectory: ../escape\n', encoding="utf-8")
    with pytest.raises(DocumentError, match="duplicate key"):
        Store(store.root)


def test_path_escape_and_existing_project_are_protected(tmp_path):
    (tmp_path / "src").mkdir()
    code = tmp_path / "src/app.py"
    code.write_text("original", encoding="utf-8")
    with pytest.raises(DocumentError):
        Store.init(tmp_path, "../outside")
    store = Store.init(tmp_path, "project-knowledge")
    assert code.read_text() == "original"
    assert not (tmp_path / "pyproject.toml").exists()
    with pytest.raises(DocumentError):
        store.document("../escape")


def test_evidence_tampering_and_post_record_edit_are_detected(project):
    store, _, proposal = project
    directory = ready(store, proposal)
    meta = read(directory / "document.yml")
    evidence = store.evidence / "extractions" / f"{meta['extraction']}.json"
    data = read(evidence)
    data["parser"] = "tampered"
    write(evidence, data)
    with pytest.raises(DocumentError, match="evidence changed"):
        store.inspect(directory)


def test_pending_writeback_is_reported_not_silently_skipped(project):
    store, _, proposal = project
    prompt = store.root / "prompt.txt"
    prompt.write_text("test", encoding="utf-8")
    store.record(proposal.name, "human", "editor", prompt)
    store.adopt(proposal.name, "reviewer")
    plan = store.export("order-design")
    assert not plan["complete"] and plan["unreflected"]
    with pytest.raises(DocumentError, match="unresolved"):
        store.export("order-design", store.root / ".arp/out/result.xlsx")


def test_cli_check_watch_schema_and_search_list(project, capsys):
    store, _, proposal = project
    ready(store, proposal)
    assert cli.main(["documents", "check", "--root", str(store.root), "--require-reviewed", "--format", "json"]) == 0
    assert '"valid": true' in capsys.readouterr().out
    assert cli.main(["documents", "watch", "--root", str(store.root), "--once"]) == 0
    assert cli.main(["documents", "list", "--root", str(store.root)]) == 0
    out = capsys.readouterr().out
    assert "content/sheet-1.md" in out and ".arp/evidence" not in out
    assert cli.main(["documents", "schema", "extraction"]) == 0


@pytest.mark.parametrize("value", ['a|b\\c\n日本語', '=SUM(A1:A2)', '<tag>& text', '_x000A_', ''])
def test_literal_strings_are_not_formulas_or_lost(tmp_path, value):
    book = Workbook()
    book.active["A1"] = "before"
    source, output = tmp_path / "source.xlsx", tmp_path / "output.xlsx"
    book.save(source)
    excel.patch(source, output, [{"sheet": "Sheet", "cell": "A1", "after": value}])
    cell = load_workbook(output).active["A1"]
    assert cell.value == value and cell.data_type == "s"


def test_markdown_import_has_structured_json_and_editable_candidate(tmp_path):
    source = tmp_path / "design.md"
    source.write_text("# 設計\n\n## 受注\n\n受注番号を保存する。\n", encoding="utf-8")
    store = Store.init(tmp_path)
    proposal = store.import_source(source, "design")
    result = store.inspect(proposal)
    assert result["extraction"]["pages"]
    assert result["pages"]


@pytest.mark.parametrize("filename", ["受注登録機能仕様書.docx", "方式提案.pptx", "検収仕様書.pdf", "得意先マスタ移行.csv"])
def test_real_office_examples_import(tmp_path, filename):
    original = Path(__file__).resolve().parents[1] / "examples/from-documents/資料" / filename
    source = tmp_path / filename
    shutil.copyfile(original, source)
    store = Store.init(tmp_path)
    proposal = store.import_source(source, "example")
    result = store.inspect(proposal)
    assert result["extraction"]["pages"]
    assert not result["reviewed"]


def test_removed_cell_requires_explicit_omission(project):
    store, _, proposal = project
    page = proposal / "content/sheet-1.md"
    page.write_text(page.read_text(encoding="utf-8").replace("| c-1-B1 | number | 10 |\n", ""), encoding="utf-8")
    mappings = read(proposal / "mappings.yml")
    mappings["entries"] = [m for m in mappings["entries"] if m["field"] != "c-1-B1"]
    write(proposal / "mappings.yml", mappings)
    with pytest.raises(DocumentError, match="source coverage"):
        store.inspect(proposal)
    mappings["omissions"] = [{"origin": None, "target": {"sheet": "項目定義", "cell": "B1"},
                              "reason": "この項目は正本から削除。原本セルの削除は今回の対象外"}]
    write(proposal / "mappings.yml", mappings)
    directory = ready(store, proposal)
    assert store.inspect(directory)["omissions"]
    assert store.export("order-design")["omissions"]


def test_damaged_formation_invalidates_review(project):
    store, _, proposal = project
    directory = ready(store, proposal)
    formation = read(directory / "formation.json")
    formation["model"] = "invented-model"
    write(directory / "formation.json", formation)
    with pytest.raises(DocumentError):
        store.inspect(directory, require_reviewed=True)


def test_record_then_edit_must_be_recorded_again(project):
    store, _, proposal = project
    prompt = store.root / "prompt.txt"
    prompt.write_text("test", encoding="utf-8")
    store.record(proposal.name, "human", "editor", prompt)
    page = proposal / "content/sheet-1.md"
    page.write_text(page.read_text(encoding="utf-8").replace("10 |", "11 |"), encoding="utf-8")
    with pytest.raises(DocumentError, match="record formation again"):
        store.adopt(proposal.name, "reviewer")


def test_schema_unknown_fields_and_version_fail(project):
    store, _, proposal = project
    meta = read(proposal / "document.yml")
    meta["schema_version"] = "2"
    write(proposal / "document.yml", meta)
    with pytest.raises(DocumentError, match="expected"):
        store.inspect(proposal)


def test_legacy_migration_and_spec_bridge(project):
    from arp4 import parse as legacy_parse
    from arp4.paths import Round
    store, source, _ = project
    round_ = Round(store.root, "r001")
    targets, _ = legacy_parse.plan(round_, [source], store.root, use_ocr=False)
    written, _ = legacy_parse.write(targets)
    text = written[0].read_text(encoding="utf-8")
    written[0].write_text(text + "\n人が追記した特別な制約。\n", encoding="utf-8")
    proposal = store.migrate(source, "migrated", round_.parsed)
    assert "人が追記した特別な制約" in (proposal / "content/page-1.md").read_text(encoding="utf-8")
    ready(store, proposal)
    result = store.prepare_spec("migrated")
    assert result["round"] == "r002" and result["files"] > 0


def test_typed_table_literals_roundtrip(tmp_path):
    values = ['a|b\\c\n日本語', '<br>', '`code`', '"quoted"', True, None, 1.25]
    page = tmp_path / "page.md"
    page.write_text('---\nschema_version: "1"\ndocument_id: d\npage_id: p\n---\n\n# Title\n\n<!-- arp:block id="b" -->\n\n'
        + '| id | type | value |\n| --- | --- | --- |\n'
        + '\n'.join(markdown.value_row(f"v{i}", value) for i, value in enumerate(values)), encoding="utf-8")
    # Typed values are raw JSON, including literal HTML; never metadata or links.
    result = markdown.parse(page, "d")
    assert list(result.values.values()) == values


def test_excel_preserves_opaque_macro_and_drawing_parts(tmp_path):
    book = Workbook()
    book.active["A1"] = "before"
    source, output = tmp_path / "macro.xlsm", tmp_path / "output.xlsm"
    book.save(source)
    # Opaque payloads exercise preservation independently of parser understanding.
    with ZipFile(source, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"opaque macro payload\x00\xff")
        archive.writestr("xl/drawings/customShape.xml", b'<opaque drawing="preserve"/>')
        archive.writestr("customXml/item1.xml", b'<vendor data="preserve"/>')
    report = excel.patch(source, output, [{"sheet": "Sheet", "cell": "A1", "after": "after"}])
    with ZipFile(source) as old, ZipFile(output) as new:
        for name in old.namelist():
            if name not in report["changed_parts"]:
                assert old.read(name) == new.read(name)


def test_signed_excel_is_rejected(tmp_path):
    book = Workbook()
    book.active["A1"] = "before"
    source, output = tmp_path / "signed.xlsx", tmp_path / "output.xlsx"
    book.save(source)
    with ZipFile(source, "a") as archive:
        archive.writestr("_xmlsignatures/sig1.xml", b"signature")
    with pytest.raises(DocumentError, match="signature"):
        excel.patch(source, output, [{"sheet": "Sheet", "cell": "A1", "after": "after"}])
    assert not output.exists()


def test_unclosed_fence_rejected(tmp_path):
    page = tmp_path / "page.md"
    page.write_text('---\nschema_version: "1"\ndocument_id: d\npage_id: p\n---\n\n# Title\n\n'
                    '<!-- arp:block id="b" -->\n\n```text\nunfinished\n', encoding="utf-8")
    with pytest.raises(DocumentError, match="unclosed"):
        markdown.parse(page, "d")
