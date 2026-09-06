"""① 切れ目の疑い ―― **黙って割れたことを、機械が言えるところまで言う。**

`suspect` が出すのは疑いであって誤りではない。だから検体は 2 つ 1 組で置く ――
**割れているものに出る**ことと、**割れていないものに出ない**ことを両方見ないと、
「全部に出る検査」でもテストは緑になる。
"""

from __future__ import annotations

from openpyxl import Workbook

from arp4 import parse, suspect
from arp4.paths import Paths, Round
from conftest import codes, parsed, sources_dir


def _sheet(*chunks: str) -> str:
    """1 シートぶんのパース結果。**塊だけを並べる。**"""
    head = "# a.xlsx / 一覧\n\n<!-- source: 資料/a.xlsx / シート: 一覧 -->\n"
    return head + "".join(chunks)


def _table(anchor: str, at: str, rows: list[list[str]]) -> str:
    """格子 1 つ。**区切り行は 1 本目の下**（:func:`arp4.mdio.dump` と同じ形）。"""
    lines = ["| " + " | ".join(row) + " |" for row in rows]
    lines.insert(1, "|" + "---|" * len(rows[0]))
    return (f"\n## 表 {at}  <!-- a:{anchor} at={at} -->\n\n"
            + "\n".join(lines) + "\n")


def _listed(anchor: str, at: str, cells: list[tuple[str, str]]) -> str:
    body = "".join(f"- `{ref}` {value}\n" for ref, value in cells)
    return f"\n## セル {at}  <!-- a:{anchor} at={at} -->\n\n{body}\n"


def _look(round_: Round, name: str, body: str):
    parsed(round_, name, body)
    findings, files, chunks = suspect.look(round_)
    return findings, files, chunks


# ── S001 方眼紙の横結合で砕けた残骸 ────────────────────────────
def test_幅1の塊が横に並んだらS001(project: Paths, round_: Round) -> None:
    """**資料に隙間は無く、機械が自分で作っていた**ぶんの残骸である。

    値が結合の左上にしか無いので、論理列 1 本ずつが幅 1 の塊になる ―― 行の
    範囲は 1 行も違わない（同じ表の同じ行だったのだから当たり前である）。
    """
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:C9", [["列1", "列2", "列3"]] + [["v", "v", "v"]] * 8),
        _listed("s1-x1", "G1:G9", [("G1", "列4")] + [(f"G{r}", "v") for r in range(2, 10)]),
        _listed("s1-x2", "K1:K9", [("K1", "列5")] + [(f"K{r}", "v") for r in range(2, 10)])))

    assert codes(findings) == ["S001"]
    assert "s1-x1 s1-x2" in findings[0].target
    assert "撮り直す" in (findings[0].hint or "")


def test_撮り直したあとはS001が消える(project: Paths, round_: Round) -> None:
    """**直ったことを機械が言えるところまでを 1 組にする。**

    検査だけあって「直った」が言えないと、エージェントは自分の直しを自分で
    採点することになる ―― 通しの中で信じられるのは、同じ機械が出す 0 件だけ
    である。ここは実物と同じ経路（`parse` → `suspect`）を通す。
    """
    path = sources_dir(project) / "方眼紙.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "一覧"
    for index in range(3):                           # 論理列 1 本を 3 列で結合
        left = 1 + index * 3
        sheet.cell(row=1, column=left, value=f"列{index + 1}")
        for row in range(2, 10):
            sheet.cell(row=row, column=left, value="v")
        for row in range(1, 10):
            sheet.merge_cells(start_row=row, start_column=left,
                              end_row=row, end_column=left + 2)
    book.save(path)

    targets, _ = parse.plan(round_, [path], sources_dir(project))
    parse.write(targets)
    findings, files, _ = suspect.look(round_)

    assert files == 1
    assert findings == []


# ── S002 ページごとに切られた 1 つの表 ────────────────────────
def test_見出しが繰り返された縦並びはS002(project: Paths, round_: Round) -> None:
    """ページの変わり目に空行が 2 行あると、機械にはそこが切れ目に見える。"""
    head = ["項番", "カラム名", "型"]
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "B1:D5", [head] + [[str(n), f"c{n}", "V"] for n in (1, 2, 3, 4)]),
        _table("s1-t2", "B8:D12", [head] + [[str(n), f"c{n}", "V"] for n in (5, 6, 7, 8)])))

    assert codes(findings) == ["S002"]
    assert "項番 | カラム名 | 型" in findings[0].message


def test_3ページに割れても1件で報せる(project: Paths, round_: Round) -> None:
    """対ごとに出していたころ、6 ページの表は 5 件になっていた ―― **同じ 1 つの
    ことを 5 回言われると、読み手は件数で重さを測れなくなる。**"""
    head = ["項番", "名前"]
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:B3", [head, ["1", "a"], ["2", "b"]]),
        _table("s1-t2", "A6:B8", [head, ["3", "c"], ["4", "d"]]),
        _table("s1-t3", "A11:B13", [head, ["5", "e"], ["6", "f"]])))

    assert codes(findings) == ["S002"]
    assert "s1-t1 s1-t2 s1-t3" == findings[0].target


def test_表題の行だけが同じでも黙る(project: Paths, round_: Round) -> None:
    """1 列しか埋まっていない行は**表題**であって見出しではない ―― それが
    一致しても「同じ表題の別の表」でしかない。"""
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:B3", [["受注一覧", ""], ["1", "a"], ["2", "b"]]),
        _table("s1-t2", "A6:B8", [["受注一覧", ""], ["x", "y"], ["z", "w"]])))

    assert findings == []


def test_見出しが違う縦並びには出さない(project: Paths, round_: Round) -> None:
    """**縦に 2 つ表があるのは普通のことである。** 疑うのは繰り返しがあるときだけ。"""
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:B3", [["項番", "名前"], ["1", "a"], ["2", "b"]]),
        _table("s1-t2", "A6:B8", [["画面ID", "画面名"], ["S1", "受注"], ["S2", "出荷"]])))

    assert findings == []


# ── S003 左右に切られた 1 つの表 ──────────────────────────────
def test_行の範囲が同じ塊が横に並んだらS003(project: Paths, round_: Round) -> None:
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:B4", [["項番", "名前"], ["1", "a"], ["2", "b"], ["3", "c"]]),
        _table("s1-t2", "E1:F4", [["項番", "説明"], ["1", "x"], ["2", "y"], ["3", "z"]])))

    assert codes(findings) == ["S003"]
    assert "撮り直しでは変わりません" in (findings[0].hint or "")


def test_行の範囲が違えば横に並んでいても黙る(project: Paths, round_: Round) -> None:
    """**ぴったり一致だけを見る。** 重なり具合で緩めると、方眼紙のシートは
    どこもかしこも重なっているので全部が疑いになる。"""
    findings, _, _ = _look(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:B4", [["項番", "名前"], ["1", "a"], ["2", "b"], ["3", "c"]]),
        _table("s1-t2", "E1:F3", [["項番", "説明"], ["1", "x"], ["2", "y"]])))

    assert findings == []


# ── 掛ける先 ────────────────────────────────────────────────────
def test_番地を持たない塊には掛けない(project: Paths, round_: Round) -> None:
    """PDF のページ・Word の表・コードの骨格は番地を持たない ―― **範囲として
    読めないものを無理に読むと、資料に無い位置を機械が作る。**"""
    body = ("# 仕様書.pdf / 1 適用範囲\n\n"
            "<!-- source: 資料/仕様書.pdf / しおり: 1 適用範囲 -->\n"
            "\n## 3 ページ  <!-- a:p3-x1 at=p.3 -->\n\n本文\n"
            "\n## 4 ページ  <!-- a:p4-x1 at=p.4 -->\n\n本文\n")
    findings, files, chunks = _look(round_, "仕様書.pdf/01_適用範囲.md", body)

    assert findings == []
    assert (files, chunks) == (1, 2)


def test_絞り込むと見た本数も減ることを数で言う(project: Paths, round_: Round) -> None:
    """0 件が「無い」なのか「見ていない」なのかを、読み手が区別できるようにする
    ―― 数えないと、絞り間違えたときにも「疑いはありません」と出る。"""
    parsed(round_, "a.xlsx/一覧.md", _sheet(
        _table("s1-t1", "A1:B4", [["項番", "名前"], ["1", "a"], ["2", "b"], ["3", "c"]]),
        _table("s1-t2", "E1:F4", [["項番", "説明"], ["1", "x"], ["2", "y"], ["3", "z"]])))

    hit, files, _ = suspect.look(round_, "a.xlsx")
    miss, empty, _ = suspect.look(round_, "b.xlsx")

    assert (codes(hit), files) == (["S003"], 1)
    assert (miss, empty) == ([], 0)
